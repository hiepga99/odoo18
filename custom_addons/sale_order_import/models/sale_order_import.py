# -*- coding: utf-8 -*-

import base64
from datetime import date as dt_date, datetime
from io import BytesIO, StringIO
from itertools import groupby
from logging import getLogger

import openpyxl
import pdfplumber
from pypdf import PdfReader

from pdfminer.converter import TextConverter
from pdfminer.layout import LAParams
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.pdfpage import PDFPage
from pdfminer.pdfparser import PDFParser

from odoo import api, models, fields, _
from odoo.exceptions import ValidationError, UserError
from odoo.tools.misc import flatten

logger = getLogger(__name__)


def convert_date_format(date):
    if not date:
        return False
    if isinstance(date, datetime):
        return date.date()
    if isinstance(date, dt_date):
        return date
    try:
        return datetime.strptime(str(date).strip(), '%d/%m/%Y').date()
    except ValueError:
        return False


def convert_string(text):
    if not text:
        return ''
    if isinstance(text, float) and text.is_integer():
        return str(int(text))
    return str(text).strip()


class _Sheet:
    """xlrd-like view (nrows, cell_value) over a list of rows so the parsers read xls and xlsx alike."""

    def __init__(self, rows):
        self.rows = rows
        self.nrows = len(rows)

    def cell_value(self, row, col):
        line = self.rows[row]
        return line[col] if col < len(line) else None


def load_sheet(data, filename):
    """First worksheet of an xlsx (openpyxl) or legacy xls (xlrd) file; date cells come back as datetime."""
    if filename.lower().endswith('.xls'):
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode)
                           if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
            rows.append(row)
        return _Sheet(rows)
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    return _Sheet([list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)])


IMPORT_TYPE = {
    '2_bigC': 'pdf,PDF',
    'bach_hoa_xanh': 'pdf,PDF',
    'coopmart': 'pdf,PDF,txt,TXT',
    'medicare': 'pdf,PDF',
    'satra': 'pdf,PDF',
    'winmart': 'xls,xlsx',
    'lotte': 'xls,xlsx',
    'excel': 'xls,xlsx',
    'brg': 'pdf,PDF',
}

IMPORT_TYPE_SELECTION = [
    ('2_bigC', 'Big C'),
    ('bach_hoa_xanh', 'Bach Hoa Xanh'),
    ('coopmart', 'COOPMART'),
    ('medicare', 'MEDICARE'),
    ('winmart', 'Winmart'),
    ('lotte', 'Lotte'),
    ('satra', 'SATRA'),
    ('excel', 'Excel'),
    ('brg', 'BRG'),
]


class SaleOrderImport(models.Model):
    _name = 'sale.order.import'
    _description = 'Sale Order Import'
    _order = 'id desc'

    name = fields.Char('Name', required=True)
    date_confirmed = fields.Date('Confirm Date', copy=False)
    date_create = fields.Date('Create Date', default=fields.Date.today)
    user_id = fields.Many2one('res.users', 'User Create', default=lambda self: self.env.user)
    company_id = fields.Many2one('res.company', default=lambda self: self.env.company)
    state = fields.Selection([('draft', 'Draft'), ('confirmed', 'Confirmed'), ('cancel', 'Cancel')],
                             string='Status', required=True, readonly=True, copy=False, default='draft')
    sale_order_ids = fields.Many2many('sale.order', string='Sale Orders', copy=False)
    so_count = fields.Integer(compute='_compute_so_count', copy=False)
    line_ids = fields.One2many('sale.order.import.line', 'import_id', copy=False,
                               string='Import Sale Order')
    import_type = fields.Selection(IMPORT_TYPE_SELECTION,
                                   default='2_bigC',
                                   string='Import Type', required=True)
    # ponytail: izi restricted this to partners having partner.product.code mappings; any partner is allowed here
    partner_id = fields.Many2one('res.partner', 'Partner System')

    create_sale_order_type = fields.Selection(
        [('one2one', 'One to one'),
         ('all2one', 'All to one')],
        default='one2one', string='Create Sale Order Type', required=True)

    attachment_ids = fields.Many2many('ir.attachment', string='Files')

    file_name = fields.Char(string='File name')
    file_upload = fields.Binary(string='File upload')
    is_order_consignment = fields.Boolean('Is order consignment', default=False)

    @api.depends('sale_order_ids')
    def _compute_so_count(self):
        for item in self:
            item.so_count = len(item.sale_order_ids)

    def update_res_id_attachment(self):
        # many2many_binary uploads with res_id=0; own them so they are not orphaned
        for record in self:
            record.attachment_ids.sudo().write({'res_model': self._name, 'res_id': record.id})

    @api.model_create_multi
    def create(self, vals_list):
        records = super().create(vals_list)
        records.update_res_id_attachment()
        return records

    def write(self, vals):
        res = super(SaleOrderImport, self).write(vals)
        if 'attachment_ids' in vals:
            self.update_res_id_attachment()
        return res

    def unlink(self):
        for rec in self:
            if rec.state not in ['draft', 'cancel']:
                raise ValidationError(_("You cannot delete record state that not be Draft or Cancel."))
        return super(SaleOrderImport, self).unlink()

    def action_confirm(self):
        self.ensure_one()
        if not self.line_ids:
            raise ValidationError(_("No data in import SO lines"))
        if self.state != 'draft':
            return
        if self.import_type == 'excel':
            order_lists = self._prepare_order_for_import_excel()
        elif self.create_sale_order_type == 'all2one':
            order_lists = [self._prepare_one_order_for_all_file(partner=self.partner_id)]
        else:
            order_lists = self._prepare_one_order_for_each_file()
        # name/price_unit/tax of each line are precomputed by sale.order.line unless given in vals
        orders = self.env['sale.order'].create(order_lists)
        self.write({
            'date_confirmed': fields.Date.today(),
            'state': 'confirmed',
            'sale_order_ids': [(6, 0, orders.ids)],
        })
        return self.action_view_sale_order()

    def _prepare_sale_order_value(self, partner, start_line):
        # ponytail: izi branch pricelist / consignment flag dropped; the partner pricelist is used
        value = {
            'partner_id': partner.id,
            'partner_invoice_id': partner.id,
            'user_id': partner.user_id.id,
            'payment_term_id': partner.property_payment_term_id.id,
            'pricelist_id': partner.property_product_pricelist.id,
            'warehouse_id': start_line.warehouse_id.id,
        }
        if start_line.po_name:
            value['client_order_ref'] = start_line.po_name
        if start_line.delivery_date:
            value['commitment_date'] = start_line.delivery_date
        return value

    def _prepare_sale_order_line_value(self, line, quantity):
        value = {
            'product_id': line.product_id.id,
            'product_uom_qty': quantity,
            'tax_id': line.tax_ids.ids,
        }
        if line.price_unit:
            value['price_unit'] = line.price_unit
        if line.discount:
            value['discount'] = line.discount
        return value

    def _prepare_one_order_for_all_file(self, lines=None, partner=None):
        lines = lines or self.line_ids
        start_line = lines[0]
        partner = partner or start_line.partner_id
        order_vals = self._prepare_sale_order_value(partner, start_line)
        quantities, first_line = {}, {}
        for line in lines:
            quantities[line.product_id] = quantities.get(line.product_id, 0) + line.quantity
            first_line.setdefault(line.product_id, line)
        order_vals['order_line'] = [
            (0, 0, self._prepare_sale_order_line_value(first_line[product], qty))
            for product, qty in quantities.items()
        ]
        return order_vals

    def _group_lines(self, key):
        """Split self.line_ids into recordsets sharing the same key(line)."""
        return [self.env['sale.order.import.line'].concat(*group)
                for _value, group in groupby(self.line_ids.sorted(key), key=key)]

    def _prepare_one_order_for_each_file(self):
        if any(self.line_ids.mapped('po_location_code')):
            groups = self._group_lines(lambda line: (line.po_location_code or '', line.po_name or ''))
        else:
            groups = self._group_lines(lambda line: line.attachment_id.id)
        return [self._prepare_one_order_for_all_file(lines) for lines in groups]

    def _find_products(self, codes):
        """{code: product} matched on default_code first, then barcode, in two queries."""
        codes = {code for code in codes if code}
        Product = self.env['product.product']
        found = {}
        for product in Product.search([('default_code', 'in', list(codes))]):
            found.setdefault(product.default_code, product)
        missing = list(codes - set(found))
        if missing:
            for product in Product.search([('barcode', 'in', missing)]):
                found.setdefault(product.barcode, product)
        return found

    def _find_partner(self, po_code):
        """Customer for a store/PO code, inside the import partner's company tree when one is set."""
        domain = [('type', 'not in', ('invoice', 'delivery'))]
        if self.import_type != 'excel' and self.partner_id:
            domain.append(('id', 'child_of', self.partner_id.id))
        if self.import_type == 'satra':
            # ponytail: izi SIMILARITY(street) lookup replaced by ilike; switch to pg_trgm if addresses drift
            domain.append(('street', 'ilike', po_code.strip()))
        else:
            # ponytail: izi matched x_po_store before ref; only ref exists here
            domain.append(('ref', '=', po_code))
        return self.env['res.partner'].search(domain, limit=1)

    def _prepare_import_order_line(self, po_code, list_data):
        records, message = [], ''
        if self.import_type == 'excel':
            codes = [str(data.get('product_code')) for data in list_data]
        elif self.import_type in ('bach_hoa_xanh', 'satra', 'brg'):
            codes = [str(data[0]) for data in list_data] + [str(data[1]) for data in list_data]
        else:
            codes = [str(data[0]) for data in list_data]
        products = self._find_products(codes)

        partner, po_name = self.partner_id, ''
        if po_code:
            if self.import_type == 'coopmart':
                po_code, po_name = po_code[0], po_code[1]
            partner = self._find_partner(po_code)
            if not partner:
                return [], '\n\t - Thiếu %s "%s" trên hệ thống.' % (
                    'địa chỉ đối tác' if self.import_type == 'satra' else 'mã đối tác', po_code)

        for data in list_data:
            product_code = str(data.get('product_code')) if self.import_type == 'excel' else str(data[0])
            product = products.get(product_code)
            if not product and self.import_type in ('bach_hoa_xanh', 'satra', 'brg'):
                product_code = str(data[1])
                product = products.get(product_code)
            if not product:
                message += '\n\t - Thiếu mã sản phẩm/barcode "%s" trên hệ thống.' % product_code
                continue
            vals = {
                'partner_id': partner.id,
                'product_id': product.id,
                'product_code': product_code,
                'po_location_code': partner.ref if self.import_type in ('satra', 'excel') else po_code,
                'po_name': po_name,
                'description': '',
                'so_number': '',
            }
            if self.import_type == 'excel':
                vals.update({
                    'quantity': data.get('product_qty'),
                    'po_name': data.get('po_name') or '',
                    'description': data.get('description') or '',
                    'so_number': data.get('so_number') or '',
                    'order_date': data.get('order_date') or False,
                    'delivery_date': data.get('delivery_date') or False,
                    'price_unit': data.get('price_unit') or 0,
                    'discount': data.get('discount') or 0,
                })
            elif self.import_type in ('bach_hoa_xanh', 'satra', 'brg'):
                vals['quantity'] = data[2]
                if self.import_type == 'brg':
                    vals['po_name'] = data[3]
            else:
                vals['quantity'] = data[1]
                if self.import_type == '2_bigC':
                    vals['po_name'] = data[2]
                elif self.import_type == 'winmart':
                    vals['po_name'] = data[2].get('po_name')
                    vals['delivery_date'] = data[2].get('commitment_date') or False
            records.append((0, 0, vals))
        return records, message

    def _generate_sale_order_data(self, value, attachment):
        list_data = []
        po_code = ''
        action = '_get_data_%s' % self.import_type
        if hasattr(self, action):
            func = getattr(self, action)
            if self.import_type in ('lotte', 'brg'):
                po_code, list_data, message = func(value)
                if message:
                    return [], message
            elif self.import_type in ('winmart', 'excel'):
                list_data, message = func(value)
                if message:
                    return [], message
            elif self.import_type in ('2_bigC', 'bach_hoa_xanh'):
                po_code, list_data = func(value)
            elif self.import_type == 'coopmart':
                try:
                    list_data = func(value)
                except UserError as e:
                    raise UserError(_("Import file error: %s", e))
                except Exception:
                    # second text engine for layouts pdfminer cannot split
                    reader = PdfReader(BytesIO(base64.b64decode(attachment.datas)))
                    list_data = func(''.join(page.extract_text() or '' for page in reader.pages))
            else:
                list_data = func(value)

        records = []
        messages = ''
        if self.import_type in ('coopmart', 'winmart'):
            for value in list_data:
                record, message = self._prepare_import_order_line(value[0], [value[1]])
                records += record
                messages += message
        elif self.import_type in ('satra', 'excel'):
            for value in list_data:
                record, message = self._prepare_import_order_line(value[0], value[1])
                records += record
                messages += message
        else:
            records, messages = self._prepare_import_order_line(po_code, list_data)

        return records, messages

    def check_type_attachment(self):
        import_type = IMPORT_TYPE.get(self.import_type, '')
        if not import_type:
            raise ValidationError("Không tìm thấy loại file hỗ trợ cho Loại import đã chọn. "
                                  "Vui lòng liên hệ quản trị viên của bạn để kiểm tra lại.")
        import_type_list = import_type.split(',')
        if any(attachment.name.split('.')[-1] not in import_type_list
               for attachment in self.attachment_ids):
            message = 'Loại import đã chọn chỉ hỗ trợ đọc file ở định dạng "%s". ' \
                      'Vui lòng kiểm tra lại danh sách file import.' % import_type
            raise ValidationError(message)
        return True

    def action_read_file(self):
        self.ensure_one()
        if self.state != 'draft':
            return
        if self.line_ids:
            self.line_ids.unlink()
        if not self.attachment_ids:
            raise ValidationError(_("Please attach at least one file."))
        self.check_type_attachment()
        lines = []
        messages = ''
        for attachment in self.attachment_ids:
            value_line, message_error = self._generate_order_from_attachment(attachment)
            if message_error:
                messages += message_error
            else:
                for value in value_line:
                    value[2]['attachment_id'] = attachment.id
                lines += value_line
        if messages:
            raise ValidationError(messages)
        self.line_ids = lines

    def _read_value_from_txt(self, attachment):
        raise UserError('Định dạng TXT chưa đọc được.')

    def _generate_order_from_attachment(self, attachment):
        value = []
        if self.import_type != 'excel':
            if 'text/plain' in attachment.mimetype and self.import_type == 'coopmart':
                value = self._read_value_from_txt(attachment)
            elif 'pdf' in attachment.mimetype:
                value = self._read_value_from_pdf(attachment)
            elif attachment.name.endswith('xls') or attachment.name.endswith('xlsx'):
                value = self._read_value_from_excel(attachment)
        else:
            value = self._read_value_from_import_SO_excel(attachment)
        try:
            value_line, message = self._generate_sale_order_data(value, attachment)
        except Exception as e:
            logger.exception(e)
            raise ValidationError(e)
        if message:
            if self.import_type != 'excel':
                message = '\n Lỗi tại file "%s":%s' % (attachment.name, message)
            else:
                message = '\n Lỗi tại file "%s":%s' % (self.file_name, message)
        return value_line, message

    def _prepare_order_for_import_excel(self):
        order_lists = []
        for lines in self._group_lines(lambda line: (line.po_location_code or '', line.so_number or '')):
            order_vals = self._prepare_one_order_for_all_file(lines)
            order_vals['note'] = lines[0].description or ''
            if lines[0].order_date:
                order_vals['date_order'] = lines[0].order_date
            order_lists.append(order_vals)
        return order_lists

    def action_cancel(self):
        self.ensure_one()
        if self.state != 'confirmed':
            return
        if any(order.state not in ('draft', 'cancel') for order in self.sale_order_ids):
            raise ValidationError(_("Cannot delete Sale Orders have state not in draft, cancel!!"))
        for order in self.sale_order_ids:
            order.unlink()
        self.line_ids.unlink()
        self.state = 'cancel'

    def action_back_to_draft(self):
        if self.state == 'cancel':
            self.state = 'draft'

    def action_view_sale_order(self):
        action = self.env['ir.actions.actions']._for_xml_id('sale.action_quotations_with_onboarding')
        action['domain'] = [('id', 'in', self.sale_order_ids.ids)]
        action['context'] = {}
        return action

    def _read_value_from_pdf(self, attachments):
        try:
            content = base64.b64decode(attachments.datas)
            if self.import_type == 'satra':
                txt = ''
                txt_list = []
                with BytesIO(content) as documents:
                    temp = pdfplumber.open(documents)
                    for page in range(len(temp.pages)):
                        text = temp.pages[page].extract_text()
                        start = text.find('Ghi chú\n')
                        stop = text.find('\nTổng cộng')
                        if start != -1 and stop != -1:
                            txt += text[start: stop + len('\nTổng cộng')]
                        # Lấy địa ghỉ giao hàng nếu có
                        l_start = text.find('Địa chỉ giao hàng: ')
                        l_stop = text.find('\nĐịa chỉ thanh toán:')
                        if l_start != -1 and l_stop != -1:
                            txt += text[l_start: l_stop + len('\nĐịa chỉ thanh toán:')]
                        # lấy đến khi có địa chỉ giao hàng là hết 1 đơn hàng, chuyển qua đơn hàng khác
                        if txt and l_start != -1 and l_stop != -1:
                            txt_list.append(txt)
                            txt = ''
                return '|||||'.join(txt_list)  # trả về danh sách đơn hàng(gồm mặt hàng và địa chỉ giao hàng)
            elif self.import_type == 'brg':
                txt = ''
                text = ''
                txt_list = []
                with BytesIO(content) as documents:
                    temp = pdfplumber.open(documents)
                    for page in range(len(temp.pages)):
                        text += temp.pages[page].extract_text()
                        start = text.find('®¬n nhËn\n')
                        stop = text.find('\nTæng gi¸ trÞ ®¬n hµng')
                        if start != -1 and stop != -1:
                            txt += text[start: stop + len('\nTæng gi¸ trÞ ®¬n hµng')]
                        else:
                            continue
                        # Lấy mã po_code
                        address = text[text.find('N¬i nhËn: ') + len('N¬i nhËn: '): text.find('\n§Þa chØ:')]
                        if address:
                            txt += address[0: address.find(' ')]
                        po_name = text[text.find('Sè §¬n: ') + len('Sè §¬n: '): text.find('\n§i¹ chØ: ')]
                        if po_name:
                            txt += f"|||{po_name}"
                        # lấy đến khi có địa chỉ giao hàng là hết 1 đơn hàng, chuyển qua đơn hàng khác
                        if txt and address:
                            txt_list.append(txt)
                            txt = ''
                        text = ''
                return '|||||'.join(
                    txt_list) if txt_list else ''  # trả về danh sách đơn hàng(gồm mặt hàng và địa chỉ giao hàng)
            elif self.import_type == 'medicare':
                output_string = ''
                with BytesIO(content) as documents:
                    temp = pdfplumber.open(documents)
                    for page in range(len(temp.pages)):
                        output_string += temp.pages[page].extract_text()
                return output_string
            else:
                output_string = StringIO()
                with BytesIO(content) as documents:
                    parser = PDFParser(documents)
                    doc = PDFDocument(parser)
                    rsrcmgr = PDFResourceManager()
                    device = TextConverter(rsrcmgr, output_string, laparams=LAParams())
                    interpreter = PDFPageInterpreter(rsrcmgr, device)
                    list_pageid = []
                    for page in PDFPage.create_pages(doc):
                        if page.pageid in list_pageid:
                            continue
                        list_pageid.append(page.pageid)
                        interpreter.process_page(page)

                return output_string.getvalue()
        except Exception as e:
            logger.exception(e)
            raise ValidationError(e)

    def _get_data_bach_hoa_xanh_1(self, value):
        try:
            value_first = value.split('Mã sản phẩm/')[1:]
            list_code = []
            list_qty_1 = []
            for v_f in value_first:
                value_last = v_f.split('\nBarcode\n')[1:][0].split('\n')
                for index, v in enumerate(value_last):
                    if not v or len(v) < 5:
                        continue
                    v = v.replace(' ', '')
                    if not v.isdigit():
                        break
                    list_code.append(v)
                list_qty_value = value_last[index:]
                for v_1 in list_qty_value:
                    if not v_1 or len(v_1) < 4 or len(v_1) > 7 or ' ' in v_1:
                        continue
                    a = v_1[:1]
                    if a.isdigit():
                        list_qty_1.append(v_1)

            list_default_code = []
            list_barcode = []

            i = 0
            while i < len(list_code):
                if (i + 1) == len(list_code):
                    break
                list_default_code.append(list_code[i].strip())
                list_barcode.append(list_code[i + 1])
                i += 2

            list_qty = []
            i = 0
            while len(list_qty) < (len(list_code) / 2):
                if len(list_qty) == (len(list_code) / 2):
                    break
                list_qty.append(float(list_qty_1[i].replace(',','')))
                # if len(list_qty_1) % 3 == 0:
                #     i += 3
                # else:
                i += 1

            list_value = list(zip(list_default_code, list_barcode, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_bach_hoa_xanh_2(self, value):
        try:
            value_first = value.split('Mã sản phẩm/')[1:]
            list_code = []
            for v_f in value_first:
                value_last = v_f.split('\nBarcode\n')[1:][0].split('\n')
                for index, v in enumerate(value_last):
                    if not v or len(v) < 5:
                        continue
                    v = v.replace(' ', '')
                    if not v.isdigit():
                        break
                    list_code.append(v)
            list_default_code = []
            list_barcode = []

            i = 0
            while i < len(list_code):
                if (i + 1) == len(list_code):
                    break
                list_default_code.append(list_code[i].strip())
                list_barcode.append(list_code[i + 1])
                i += 2

            value_qty_first = value.split('\nGhi chú\n')[1:][0]
            value_qty_last = value_qty_first.split('Mã sản phẩm/')[0].split('\n')
            list_qty_1 = []
            for v in value_qty_last:
                if not v or '.' not in v:
                    continue
                list_qty_1.append(v)

            list_qty = []
            i = 0
            while len(list_qty) < (len(list_code) / 2):
                if len(list_qty) == (len(list_code) / 2):
                    break
                list_qty.append(float(list_qty_1[i].replace(',','')))
                if len(list_qty_1) % 3 == 0:
                    i += 3
                else:
                    i += 4
            list_value = list(zip(list_default_code, list_barcode, list_qty))
            return list_value
        except Exception:
            return []

    def _get_bach_hoa_xanh_po_code1(self, value):
        try:
            value_first_po_code = value.split('ĐƠN ĐẶT HÀNG')[1:][0]
            value_last_po_code = value_first_po_code.split('Tên NCC:\n')[0].split(':')
            if len(value_last_po_code) == 3:
                value_last_po_code = value_last_po_code[2].split('-')[0]
            elif len(value_last_po_code) == 1:
                value_last_po_code_1 = value_last_po_code[0].split('-')
                value_last_po_code = value_last_po_code_1[1]
            elif len(value_last_po_code) == 4:
                try:
                    value_last_po_code_4 = value_last_po_code[0].split('-')
                    value_last_po_code = value_last_po_code_4[1]
                except Exception:
                    value_last_po_code = value_last_po_code[3].split('-')
                    value_last_po_code = value_last_po_code[1]
            else:
                try:
                    value_last_po_code = value_last_po_code[3].split('-')
                    if len(value_last_po_code) == 4:
                        value_last_po_code = value_last_po_code[1]
                    else:
                        value_last_po_code = value_last_po_code[2]
                except Exception:
                    value_last_po_code = value_first_po_code.split('\nĐịa chỉ kho NCC giao:')[0].split(':')
                    value_last_po_code = value_last_po_code[3].split('-')[0]
            po_code = value_last_po_code.split('\n')[-1].strip()
            return po_code
        except Exception:
            return ''

    def _get_bach_hoa_xanh_po_code2(self, value):
        try:
            value_first_po_code = value.split('Fax: ')[1:][0]
            value_last_po_code = value_first_po_code.split('Tên NCC:\n')[0].split('\n')
            po_code = ''
            for v in value_last_po_code:
                if '-' not in v:
                    continue
                v_1 = v.split(' - ')
                if len(v_1) == 3:
                    po_code = v_1[0]
                    break
            return po_code
        except Exception:
            return ''

    def _get_data_bach_hoa_xanh(self, value):
        po_code = self._get_bach_hoa_xanh_po_code1(value)
        if not po_code.isdigit():
            po_code = self._get_bach_hoa_xanh_po_code2(value)
        list_value = self._get_data_bach_hoa_xanh_1(value)
        if not list_value:
            list_value = self._get_data_bach_hoa_xanh_2(value)

        return po_code, list_value

    def _get_data_2_bigC(self, value):
        po_name = value.split('Ad.Ch\n')[1:][0].split('\n')[1]
        try:
            value_name_first = value.split('By Supplier\n')[1:][0]
            value_name_last = value_name_first.split('Article desc.')[0].split('\n')
            delivered_to = value_name_last[1] if not value_name_last[0] else value_name_last[0]
            po_code = delivered_to.split('(')[1].split(')')[0]
        except Exception:
            try:
                value_name_first = value.split('Delivered To\n')[1:][0]
                value_name_last = value_name_first.split('Article desc.')[0].split('\n')
                delivered_to = value_name_last[1] if not value_name_last[0] else value_name_last[0]
                po_code = delivered_to.split('(')[1].split(')')[0]
            except Exception:
                value_name_first = value.split('Ordered By\n')[1:][0]
                value_name_last = value_name_first.split('Delivered To\n')[0].split('\n')
                delivered_to = ''
                for v in value_name_last:
                    if '(' in v and ')' in v:
                        delivered_to = v
                        break
                po_code = delivered_to.split('(')[1].split(')')[0]

        try:
            value_first = value.split('Article desc.')[1:][0]
            value_last = value_first.split('OU\ntype')[0].split('\n')
            list_code = []
            for r in value_last:
                if not r:
                    continue
                code = r.split(' ')[0]
                if code.isdigit():
                    list_code.append(code)
            qty_first = value.split('Free\n\nquantity')[1:][0]
            qty_last = qty_first.split('\nOrdered from')[0].split('\n')
            list_qty = []
            for q in qty_last:
                if not q:
                    continue
                if ' ' not in q:
                    q = float(q)
                    if q > 0:
                        list_qty.append(q)
            list_value = list(zip(list_code, list_qty, [po_name] * (len(list_code))))
        except Exception:
            value_first = value.split('\nArticle Desc')[1:][0]
            value_last = value_first.split('\nTOTAL BF.TAX')[0].replace(' ', '.')
            value_focus = value_last.split('OU\ntype')[0].split('\n')
            list_code = []
            list_qty = []
            for r in value_focus:
                if not r:
                    continue
                if not r.startswith('...'):
                    code = r.split('.')[0]
                    if code.isdigit() and len(code) > 4:
                        list_code.append(code)
                else:
                    qty_list = r.split('................')
                    if len(qty_list) == 3:
                        if qty_list[0]:
                            sku_ou = float(qty_list[0].replace('.', ''))
                            qty = float(qty_list[1].replace('.', ''))
                        else:
                            a = qty_list[1].split('..')
                            sku_ou = float(a[0].replace('.', ''))
                            qty = float(a[-1])
                        list_qty.append(qty * sku_ou)
                    elif len(qty_list) == 4:
                        if not qty_list[0]:
                            qty = qty_list[2].lstrip('.')
                            list_qty.append(float(qty))
            list_value = list(zip(list_code, list_qty, [po_name] * (len(list_code))))
        return po_code, list_value

    def _get_data_coopmart_bien_hoa(self, pdf):
        """Mẫu file import giống Co.opMart Bien Hoa"""
        try:
            value_code_first = pdf.split('\nSKU Number\n')[1:][0]
            list_code = value_code_first.split('\n\nDescription\n')[0].split('\n')

            value_qty_first = pdf.split('\nQty\nOrd/\nPcs\n')[1:][0]
            value_qty = value_qty_first.split('\n\nQty\nRec/\nPcs\n')[0].split('\n')

            list_qty = []
            for q in value_qty[:len(list_code)]:
                if not q:
                    continue
                list_qty.append(float(q))
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_chu_se(self, pdf):
        """Mẫu file import giống  Co.opMart Chu Se, Co.opMart Nguyen Binh"""
        try:
            value_code_first = pdf.split('\nSKU Number Description\n')[1:][0]
            value_code_last = value_code_first.split('\nSell Buy\n')[0].split('\n')

            # CM NGUYỄN BÌNH
            if value_code_last[0] == 'Cost':
                value_code_last_new = []
                for v in value_code_last[1:]:
                    if not v:
                        continue
                    v1 = v.split()
                    if len(v1) > 1:
                        value_code_last_new.append(v)
                value_code_last = value_code_last_new

            list_code = []
            for v in value_code_last:
                if not v:
                    continue
                v = v.split()[0]
                list_code.append(v)

            try:
                value_qty_first = pdf.split('\nQty\nOrd/\nPcs\n')[1:][0]
                value_qty = value_qty_first.split('\n\nQty\nRec/\nPcs\n')[0].split('\n')
            except Exception:
                # CM NGUYỄN BÌNH
                value_qty_first = pdf.split('\nQty\nOrd/\n\nPcs\n')[1:][0]
                value_qty_last = value_qty_first.split('\n\nQty\nRec/\n\nPcs\n')[0].split('\n')
                value_qty = []
                for q in value_qty_last:
                    if not q:
                        continue
                    value_qty.append(q)

            list_qty = []
            for q in value_qty[:len(list_code)]:
                if not q:
                    continue
                list_qty.append(float(q))
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_tay_ninh(self, pdf):
        """Mẫu file import giống 'Co.opMart Tay Ninh', 'Co.opMart Tra Vinh', 'Co.opMart Trang Bang'"""
        try:
            value_code_first = pdf.split('Pcs                                       Cost')[1:][0]
            value_code_last = value_code_first.split('                                    ')[0].replace(
                '\n', '').split('  ')
            list_code_value = []
            for v in value_code_last:
                v_new = v.lstrip()
                if not v or not v_new:
                    continue
                list_code_value.append(v_new)

            list_code = []
            list_qty = []
            i = 0
            while i < len(list_code_value):
                list_code.append(list_code_value[i])
                list_qty.append(float(list_code_value[i + 7]))
                i += 10
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_hai_phong(self, pdf):
        """'Co.opMart Hai Phong'"""
        try:
            value_code_first = pdf.split('Pcs                                      Cost')[1:][0]
            value_code_last = value_code_first.split(' Sub Total - ')[0].replace('\n', '').split('  ')
            list_code_value = []
            for v in value_code_last:
                v_new = v.lstrip()
                if not v or not v_new:
                    continue
                list_code_value.append(v_new)

            list_code = []
            list_qty = []
            i = 0
            while i < len(list_code_value):
                list_code.append(list_code_value[i])
                list_qty.append(float(list_code_value[i + 7]))
                i += 10
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_bao_loc(self, pdf):
        """'Co.opMart Bao Loc', 'Co.opMart Ly Thuong Kiet', 'Co.opMart Nguyen Anh Thu',
            'Co.opMart Quang Binh', 'Co.opMart Vinh Long',
            'Co.opMart Vung Tau', 'Co.opMart Xa Lo Ha Noi'
        """
        try:
            value_code_first = pdf.split('Pcs                                      Cost')[1:][0]
            value_code_last = value_code_first.split('Sub Total - ')[0].replace('\n', '').split('  ')

            list_code_value = []
            for v in value_code_last:
                v_new = v.lstrip()
                if not v or not v_new:
                    continue
                list_code_value.append(v_new)

            list_code = []
            list_qty = []
            i = 0
            while i < len(list_code_value):
                list_code.append(list_code_value[i])
                list_qty.append(float(list_code_value[i + 7]))
                i += 10
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_hoang_van_thu(self, pdf):
        """'Co.opMart Hoang Van Thu'"""
        try:
            value_code_first = pdf.split('Pcs         Pcs             Cost')[1:][0]
            value_code_last = value_code_first.split('Sub Total - ')[0].split('\xa0\xa0')
            list_code_value = []
            for v in value_code_last:
                if not v or v == ' \n':
                    continue
                list_code_value.append(v)

            list_code = []
            list_qty = []
            i = 0
            while i < len(list_code_value):
                list_code.append(list_code_value[i])
                value_qty = list_code_value[i + 7].replace('\xa0', '')
                list_qty.append(float(value_qty))
                i += 10
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_SCA_GOLDSILK(self, pdf):
        """'Co.opMart SCA - GOLDSILK'"""
        try:
            value_first = pdf.split('Vendor Part No.')[1:][0]
            value_last_code = value_first.split('\nNotes - ')[0].split('\n\n')[1].split('\n')
            list_code = []
            index_code = 0
            while index_code < len(value_last_code):
                list_code.append(value_last_code[index_code])
                index_code += 2

            value_last = value_first.split('\nSell Buy\nOrd/\n')[0].split('\n')
            list_value_qty = []
            for r in value_last:
                if not r:
                    continue
                list_value_qty.append(r)

            value_last_qty = list_value_qty[index_code:]
            list_qty = []
            index_qty = 0
            while index_qty < len(value_last_qty):
                qty = value_last_qty[index_qty + 1]
                if '.' in qty:
                    list_qty.append(qty)
                    index_qty += 3
                else:
                    qty = value_last_qty[index_qty + 2]
                    list_qty.append(qty)
                    index_qty += 4
            list_qty = [float(qty) for qty in list_qty[:len(list_code)]]
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_ha_dong(self, pdf):
        """'Co.opMart Ha Dong'"""
        try:
            value_code_first = pdf.split('     Cost\n  ')[1:][0]
            value_code_last = value_code_first.split('     Sub Total - ')[0].split('\n')
            list_code_value = []
            for v in value_code_last:
                if not v or not v.strip():
                    continue
                list_code_value.append(v.lstrip())
            list_code = []
            list_qty = []
            i = 0
            while i < len(list_code_value):
                list_code.append(list_code_value[i].split()[0])
                value_qty = list_code_value[i + 1].split('      ')[3]
                list_qty.append(float(value_qty.strip()))
                i += 3
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_ha_noi(self, pdf):
        """'Co.opMart Ha Noi'"""
        try:
            value_code_first = pdf.split('Cost\n  ')[1:][0]
            value_code_last = value_code_first.split('     Sub Total - ')[0].split('\n')
            list_code_value = []
            for v in value_code_last:
                if not v or not v.strip():
                    continue
                list_code_value.append(v.lstrip())

            list_code = []
            list_qty = []
            i = 0
            while i < len(list_code_value):
                value_split = list_code_value[i:i + 1][0].split('      ')
                list_value_split = []
                for v_1 in value_split:
                    if not v_1 or not v_1.strip():
                        continue
                    list_value_split.append(v_1.lstrip())

                code = list_value_split[0].split()[0]
                list_code.append(code)
                if len(list_value_split) == 9:
                    qty = float(list_value_split[6])
                else:
                    qty = float(list_value_split[5])
                list_qty.append(qty)
                i += 2
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_quang_ngai(self, pdf):
        """'Co.opMart SCA - VICTORIA', 'Co.opMart Quy Nhon', 'Co.opMart Quang Ngai'"""
        try:
            try:
                value_code_first = pdf.split('Cost\n  ')[1:][0]
                value_code_last = value_code_first.split('     Sub Total - ')[0].split('\n')
            except Exception:
                value_code_first = pdf.split('Cost\n\n  ')[1:][0]
                value_code_last = value_code_first.split('Sub Total -')[0].split('\n')

            list_code_value = []
            for v in value_code_last:
                if not v or not v.strip():
                    continue
                list_code_value.append(v.lstrip())

            list_code = []
            list_qty = []
            for v in list_code_value:
                value_split = v.split('     ')
                list_value_split = []
                for v_1 in value_split:
                    if not v_1 or not v_1.strip():
                        continue
                    list_value_split.append(v_1.lstrip())

                if len(list_value_split) > 4:
                    code = list_value_split[0].split()[0]
                    list_code.append(code)
                    qty = float(list_value_split[5])
                    list_qty.append(qty)
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_SCA_long_bien(self, pdf):
        """'Co.opMart SCA - LONG BIEN'"""
        try:
            value_code_first = pdf.split('\nVendor Part No. U/M\n')[1:][0]
            value_code_last = value_code_first.split('\nSell Buy\n')[0].split('\n')
            list_code = []
            for v in value_code_last:
                if not v:
                    continue
                v_new = v.split()
                if len(v_new) > 1:
                    list_code.append(v_new[0])

            value_qty_first = pdf.split('\nQty\nOrd/\nPcs\n')[1:][0]
            value_qty = value_qty_first.split('\n\nQty\nRec/\nPcs\n')[0].split('\n')
            list_qty = [float(q) for q in value_qty]
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_new(self, pdf):
        try:
            value_code_first = pdf.split('Cost\n\n')[1:][0]
            value_code_last = value_code_first.split('                               ')[0].split('        ')
            list_code_value = []
            for v in value_code_last:
                if not v or not v.strip():
                    continue
                list_code_value.append(v.lstrip())

            list_code = []
            list_qty = []
            i = 0
            while i < len(list_code_value):
                code = list_code_value[i].split('   ')[0]
                list_code.append(code)
                list_qty.append(float(list_code_value[i + 5]))
                i += 7
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_soc_trang(self, pdf):
        try:
            value_first = pdf.split('                      Cost\n')[1:][0]
            value_last = value_first.split('\nSub Total -')[0].split('\n')
            list_code_pre = []
            list_qty_pre = []
            index_code = 0
            index_qty = 2
            while index_code < len(value_last):
                list_code_pre.append(value_last[index_code])
                index_code += 4
                list_qty_pre.append(value_last[index_qty])
                index_qty += 4

            list_code = []
            for i in list_code_pre:
                if not i:
                    continue
                code = i.split('  ')[1]
                if not code:
                    continue
                list_code.append(code)

            list_qty = []
            for i in list_qty_pre:
                if not i:
                    continue
                qty_split = i.split(' ')
                qty = [i for i in qty_split if i]
                if not qty:
                    continue
                list_qty.append(float(qty[2]))
            list_qty = list_qty[:len(list_code)]
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_hoc_mon(self, pdf):
        try:
            value_first = pdf.split('Cost\n\n  ')[1:][0]
            value_last = value_first.split('Sub Total  -')[0].split('\n')[:-1]
            list_code = []
            list_qty = []
            for v in value_last:
                v = [i for i in v.split('      ') if i]
                if not v:
                    continue
                value_code = v[0].strip().split('  ')[0]
                if value_code:
                    list_code.append(value_code)
                value_qty = v[-3].strip()
                list_qty.append(float(value_qty))
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_hung_vuong(self, pdf):
        try:
            value_first = pdf.split('Cost\n\n   ')[1:][0]
            value_last = value_first.split('\n\n                      Notes -')[0].split('\n')[:-1]
            list_code = []
            list_qty = []
            for v in value_last:
                if not v:
                    continue
                v = [i for i in v.split('           ') if i]
                value_code = v[0].strip().split('  ')[0]
                if value_code:
                    list_code.append(value_code)
                value_qty = v[-3].strip()
                list_qty.append(float(value_qty))
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_tan_chau_an_giang(self, pdf):
        try:
            value_code_first = pdf.split('Vendor Part No. U/M U/M')[1:][0]
            value_code_last = value_code_first.split('\nNotes -')[0].split('\n')[:-1]
            list_code = []
            for v in value_code_last:
                if not v or ' ' not in v:
                    continue
                code = v.strip().split(' ')[0]
                list_code.append(code)

            value_qty_first = pdf.split('Qty\nOrd/\nPcs\n')[1:][0]
            value_qty_last = value_qty_first.split('Qty\nRec/\nPcs\n')[0].split('\n')[:-1]
            list_qty = [float(i) for i in value_qty_last if i]
            list_qty = list_qty[:len(list_code)]
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_cum_gar(self, pdf):
        try:
            value_code_first = pdf.split('SKU Number Description')[1:][0]
            value_code_last = value_code_first.split('\nSell Buy\n')[0].split('\n')[:-1]

            list_code = []
            for v in value_code_last:
                if not v or ' ' not in v:
                    continue
                code = v.strip().split(' ')[0]
                list_code.append(code)
            value_qty_first = pdf.split('Qty\nOrd/\nPcs\n')[1:][0]
            value_qty_last = value_qty_first.split('Qty\nRec/\nPcs\n')[0].split('\n')[:-1]
            list_qty = [float(i) for i in value_qty_last if i]
            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_head_office_SGC(self, pdf):
        try:
            try:
                value_first = pdf.split('Vendor Part No.')[1:][0]
                value_last_code = value_first.split('Sub Total -')[0].split('\n\n')[1].split('\n')
                list_code = []
                for c in value_last_code:
                    list_code.append(c.strip())

                value_last_qty = value_first.split('Qty\nRec/\nPcs\n')[1:][0].split('Sub Total -')[0].split('\n')
                list_qty_1 = []
                for q in value_last_qty:
                    if not q:
                        continue
                    list_qty_1.append(q)
                list_qty_2 = list_qty_1[:-len(list_code)]
                list_qty_3 = list_qty_2[-len(list_code):]
                list_qty = []
                for q_1 in list_qty_3:
                    list_qty.append(float(q_1.strip()))

                list_value = list(zip(list_code, list_qty))
            except Exception:
                value_first = pdf.split('Vendor Part No.')[1:][0]
                value_last_code = value_first.split('Sell Buy\n')[0].split('Qty\nRec/\nPcs\n')[1].split('\n')
                list_code = []
                for c in value_last_code:
                    if not c:
                        continue
                    if '-' not in c:
                        break
                    list_code.append(c.strip())

                value_last_qty = value_first.split('Sell Buy\n')[0].split('\n')
                list_qty_1 = []
                for q in value_last_qty:
                    if not q:
                        continue
                    list_qty_1.append(q)
                list_qty_2 = list_qty_1[:-len(list_code)]
                list_qty_3 = list_qty_2[-len(list_code):]
                list_qty = []
                for q_1 in list_qty_3:
                    list_qty.append(float(q_1.strip()))

                list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_data_coopmart_head_office_SGC_1(self, pdf):
        try:
            value_first = pdf.split('EA\n\nC30\n')[1:][0]
            value_last = value_first.split('Buy\nCost\n')[0].split('\n')
            value_last_code = []
            for l in value_last:
                if not l:
                    continue
                value_last_code.append(l)
            list_code = []
            for c in value_last_code:
                if '-' not in c:
                    break
                list_code.append(c.strip())
            list_qty_2 = value_last_code[:-len(list_code)]
            list_qty_3 = list_qty_2[-len(list_code):]
            list_qty = []
            for q_1 in list_qty_3:
                list_qty.append(float(q_1.strip()))

            list_value = list(zip(list_code, list_qty))
            return list_value
        except Exception:
            return []

    def _get_coopmart_name(self, pdf):
        po_code = ''
        coop_mart_name = ''
        po_name = ''
        try:
            po_location_value = []
            try:
                po_location_first = pdf.split('P/O Location:')[1:][0]
                po_location_last = po_location_first.split('Buyer   - ')[0].split('\n')
                for v in po_location_last:
                    if not v:
                        continue
                    v_new = v.split('    ')
                    po_location_value.append(v_new)
                po_location_value = flatten(po_location_value)
            except Exception:
                try:
                    po_location_first = pdf.split('P/O Location:\xa0')[1:][0]
                    po_location_last = po_location_first.split('Buyer   -')[0].split('\xa0\xa0\xa0')
                    for v in po_location_last:
                        if not v:
                            continue
                        for v1 in v.split(' '):
                            v_new = v1.replace('\xa0', ' ').replace('\n', '')
                            po_location_value.append(v_new)
                except Exception:
                    po_location_first = pdf.split('by SKU Number\nStore-   ')[1:][0]
                    po_location_value = po_location_first.split('               Vendor -  ')[0].split(' ')
                    po_location_value[3] = ' '.join(po_location_value[1:4])

            list_po_value = []
            for v in po_location_value:
                if not v:
                    continue
                list_po_value.append(v.lstrip())

            po_name = ''
            for v in list_po_value[:5]:
                if '-' in v and ' ' not in v:
                    po_name = v
                    break
            for index, code in enumerate(list_po_value[:5]):
                if code.isdigit():
                    po_code = code
                    if index == 0:
                        coop_mart_name = list_po_value[3]
                    else:
                        coop_mart_name = list_po_value[1]
                    break
            if not po_name:
                try:
                    po_name_first = pdf.split('P/O Number:    ')[1:][0]
                    po_name_last = po_name_first.split('Purchase Order    ')[0]
                    po_name = po_name_last.strip().split(' ')[0].split('\n')[0]
                except Exception:
                    po_name_first = pdf.split('P/O Number:    ')[1:][0]
                    po_name_last = po_name_first.split('Purchase Order      ')[0]
                    po_name = po_name_last.strip().split(' ')[0].split('\n')[0]

            if coop_mart_name in ('Head Office SGC - 1', 'CONG TY TNHH MTV MARFOUR', 'Head Office Cty TNHHMTV TP SGC'):
                coop_mart_first = pdf.split('\nShip To:')[1:][0]
                coop_mart_last = coop_mart_first.split('\nSKU Number')[0].split('\n')[1]
                coop_mart_name = coop_mart_last.split('  ')[0]
        except Exception:
            po_code = pdf.split('\n')[0].strip()
            po_location_first = pdf.split('Bi l l   To:')[1:][0]
            po_location_last = po_location_first.split('Buy er       -')[0].split('\n')
            coop_mart_name = po_location_last[1].split('             ')[0]
        return po_code, coop_mart_name, po_name

    def _get_data_pdf_coopmart(self, pdf):
        po_name = ''
        try:
            po_code, coop_mart_name, po_name = self._get_coopmart_name(pdf)
        except Exception:
            value_code_first = pdf.split('\nStore-\n')[1:][0]
            value_code_last = value_code_first.split('\nVendor -')[0].split('\n')[1]
            value_code_last = value_code_last.split(' ', 1)
            po_code, coop_mart_name = value_code_last[0], value_code_last[1]
        if coop_mart_name in ('Co.opMart Bien Hoa',):
            list_value = self._get_data_coopmart_bien_hoa(pdf)
        elif coop_mart_name in ('Co.opMart Chu Se', 'Co.opMart Nguyen Binh', 'Co.opMart Buon Ma Thuot',
                                'Co.opMart Phu Lam', 'Co.opMart Bac Giang'):
            list_value = self._get_data_coopmart_chu_se(pdf)
        elif coop_mart_name in ('Co.opMart Tay Ninh', 'Co.opMart Tra Vinh', 'Co.opMart Trang Bang'):
            list_value = self._get_data_coopmart_tay_ninh(pdf)
        elif coop_mart_name == 'Co.opMart Hai Phong':
            list_value = self._get_data_coopmart_hai_phong(pdf)

        elif coop_mart_name in ('Co.opMart Bao Loc', 'Co.opMart Ly Thuong Kiet', 'Co.opMart Nguyen Anh Thu',
                                'Co.opMart Quang Binh', 'Co.opMart Vinh Long',
                                'Co.opMart Vung Tau', 'Co.opMart Xa Lo Ha Noi'):
            list_value = self._get_data_coopmart_bao_loc(pdf)

        elif coop_mart_name == 'Co.opMart Hoang Van Thu':
            list_value = self._get_data_coopmart_hoang_van_thu(pdf)
        elif coop_mart_name in ('Co.opMart SCA - GOLDSILK',):
            list_value = self._get_data_coopmart_SCA_GOLDSILK(pdf)
        elif coop_mart_name == 'Co.opMart Ha Dong':
            list_value = self._get_data_coopmart_ha_dong(pdf)
        elif coop_mart_name in ('Co.opMart Ha Noi',):
            list_value = self._get_data_coopmart_ha_noi(pdf)
        elif coop_mart_name in ('Co.opMart SCA - VICTORIA', 'Co.opMart Quang Ngai'):
            list_value = self._get_data_coopmart_quang_ngai(pdf)
        # elif coop_mart_name in ('Head Office Cty TNHHMTV TP SGC',):
        #     list_value = self._get_data_coopmart_head_office_SGC(pdf)
        elif coop_mart_name == 'Co.opMart SCA - LONG BIEN':
            list_value = self._get_data_coopmart_SCA_long_bien(pdf)
        elif coop_mart_name == 'Co.opMart Hoc Mon':
            list_value = self._get_data_coopmart_hoc_mon(pdf)
        elif coop_mart_name == 'Co.opMart  96 Hung Vuong':
            list_value = self._get_data_coopmart_hung_vuong(pdf)
        # elif coop_mart_name == 'Co.opMart Tan Chau An Giang':
        #     list_value = self._get_data_coopmart_tan_chau_an_giang(pdf)
        elif coop_mart_name == "Co.opMart Cu M'gar":
            list_value = self._get_data_coopmart_cum_gar(pdf)
        else:
            # Đọc từng loại một
            list_value = self._get_data_coopmart_bien_hoa(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_chu_se(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_tay_ninh(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_hai_phong(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_bao_loc(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_hoang_van_thu(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_SCA_GOLDSILK(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_ha_dong(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_ha_noi(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_quang_ngai(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_SCA_long_bien(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_new(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_soc_trang(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_hoc_mon(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_hung_vuong(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_tan_chau_an_giang(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_cum_gar(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_head_office_SGC(pdf)
            if not list_value:
                list_value = self._get_data_coopmart_head_office_SGC_1(pdf)

            if not list_value:
                raise UserError('Mẫu file "%s" chưa được định nghĩa để đọc.'
                                'Vui lòng liên hệ quản trị viên để được hỗ trợ.' % coop_mart_name)
        return [po_code, po_name], list_value

    def _get_data_coopmart(self, value):
        list_value_pdf = value.split('POM343')
        list_value_multi = []
        if len(list_value_pdf) > 3:
            for pdf in list_value_pdf:
                if not pdf:
                    continue
                po_infor, list_value = self._get_data_pdf_coopmart(pdf)
                list_value_multi += list(zip([po_infor] * len(list_value), list_value))
        else:
            if len(list_value_pdf) == 3:
                pdf = list_value_pdf[2]
            elif len(list_value_pdf) == 1:
                pdf = value
            else:
                pdf = list_value_pdf[1]

            po_infor, list_value = self._get_data_pdf_coopmart(pdf)
            list_value_multi += list(zip([po_infor] * (len(list_value)), list_value))
        return list_value_multi

    def _get_data_medicare1(self, value):
        value_first = value.split('\nThanh tien\n')[1:][0]
        value_last = value_first.split('\nGhi chu:\n')[0].split('\n')
        list_code = []
        list_qty = []
        for v in value_last:
            if not v or ',' in v:
                continue
            if v.isdigit():
                list_qty.append(float(v))
            else:
                if ' ' in v:
                    v1 = v.split(' ')
                    if len(v1) == 2 and v1[1]:
                        list_code.append(v1[1])
        list_qty = list_qty[:len(list_code)]
        list_value = list(zip(list_code, list_qty))
        if not list_value:
            value_code_first = value.split('\nDien giai\n')[1:][0]
            value_code_last = value_code_first.split('\nDon vi\ntinh\n')[0].split('\n\n \n\n ')[0].split('\n')
            list_code = []
            for v in value_code_last:
                if not v:
                    continue
                code = v.split(' ')[1]
                list_code.append(code)

            value_qty_first = value.split('\nThanh tien\n')[1:][0]
            value_qty_last = value_qty_first.split('\nGhi chu:\n')[0].split('\n')
            list_qty_str = []
            for q in value_qty_last:
                if not q:
                    continue
                list_qty_str.append(q)
            list_qty_str = list_qty_str[:len(list_code)]
            list_qty = [float(q1) for q1 in list_qty_str]
            list_value = list(zip(list_code, list_qty))
        return list_value

    def _get_data_medicare(self, value):
        value_first = value.split('Medicare nha cung cap tinh luong gia\n')[1:][0]
        value_last = value_first.split('\nTong cong')[0].split('\n')
        list_code = []
        list_qty = []
        for v in value_last:
            if not v:
                continue
            v1 = v.split(' ')
            list_code.append(v1[1])
            if ',' in v1[-3]:
                list_qty.append(float(v1[-2]))
            else:
                list_qty.append(float(v1[-3]))
        list_value = list(zip(list_code, list_qty))
        return list_value

    def _get_data_winmart(self, sheet):
        rows = 1
        value = []
        message = ''
        while rows < sheet.nrows:
            try:
                po_name = convert_string(sheet.cell_value(rows, 1))
                po_code = convert_string(sheet.cell_value(rows, 9))
                product_code = convert_string(sheet.cell_value(rows, 16))
                commitment_date = convert_date_format(sheet.cell_value(rows, 24))
                extra_data = {
                    'commitment_date': commitment_date,
                    'po_name': po_name or ''
                }
                if not product_code:
                    message += f"\n\t - Thiếu mã sản phẩm tại hàng {rows + 1} trên hệ thống."
                product_qty = sheet.cell_value(rows, 19)
                if not product_qty:
                    message += f"\n\t - Không xác định số lượng tại hàng {rows + 1}."
                if product_code and product_qty:
                    value.append((po_code, [product_code, float(product_qty), extra_data]))
            except Exception as e:
                message += f"\n\t - Lỗi định dạng mã sản phẩm/số lượng tại hàng {rows + 1}." \
                           f"\nChi tiết: {str(e)}"
            rows += 1
        return value, message

    def _get_data_lotte(self, sheet):
        rows = 5
        value = []
        message = ''
        po_code = ''
        while rows < sheet.nrows:
            try:
                product_code = convert_string(sheet.cell_value(rows, 2))
                if not product_code:
                    message += f"\n\t - Thiếu mã sản phẩm/barcode tại hàng {rows + 1}."
                product_qty = sheet.cell_value(rows, 4)
                if not product_qty:
                    message += f"\n\t - Không xác định số lượng tại hàng {rows + 1}."
                if product_code and product_qty:
                    value.append((product_code, float(product_qty)))
            except Exception as e:
                message += f"\n\t - Lỗi định dạng mã sản phẩm/số lượng tại hàng {rows + 1}." \
                           f"\n\tChi tiết: {str(e)}"
            rows += 1
        return po_code, value, message

    def _get_data_excel(self, sheet):
        list_data = []
        rows = 1
        value = []
        message = ''
        po_code = ''
        po_code_temp = ''
        po_name = ''
        so_number = ''
        order_date = None
        delivery_date = None
        description = ''
        price_unit = 0
        discount = 0
        cur_so_code = False
        while rows < sheet.nrows:
            try:
                so_code = convert_string(sheet.cell_value(rows, 0))
                if so_code:
                    po_code_temp = convert_string(sheet.cell_value(rows, 1))
                    if not po_code_temp:
                        message += f"\n\t - Thiếu mã khách hàng tại hàng {rows + 1}."
                        rows += 1
                        continue
                if not po_code_temp:
                    rows += 1
                    continue
                if cur_so_code != so_code and cur_so_code and so_code:
                    cur_so_code = so_code
                    list_data.append([po_code, value])
                    value = []
                    po_name = ''
                    order_date = None
                    delivery_date = None
                    description = ''
                    price_unit = 0
                    discount = 0
                po_code = po_code_temp
                po_name = convert_string(sheet.cell_value(rows, 2)) if sheet.cell_value(rows, 2) else po_name
                order_date = convert_date_format(sheet.cell_value(rows, 3)) if sheet.cell_value(rows, 3) else order_date
                delivery_date = convert_date_format(sheet.cell_value(rows, 4)) if sheet.cell_value(rows,
                                                                                                   4) else delivery_date
                if self.is_order_consignment:
                    try:
                        price_unit = float(sheet.cell_value(rows, 9)) if sheet.cell_value(rows, 9) else 0
                        discount = float(sheet.cell_value(rows, 10)) if sheet.cell_value(rows, 10) else 0
                    except Exception as e:
                        message += f"\n\t - Lỗi định dạng đơn giá/chiết khấu tại hàng {rows + 1}." \
                                   f"\n\tChi tiết: {str(e)}"
                so_number = so_code if so_code else so_number
                description = convert_string(sheet.cell_value(rows, 5)) if sheet.cell_value(rows, 5) else description
                if not cur_so_code and so_code:
                    cur_so_code = so_code
                product_code = convert_string(sheet.cell_value(rows, 6))
                if not product_code:
                    message += f"\n\t - Thiếu mã sản phẩm/barcode tại hàng {rows + 1}"
                product_qty = sheet.cell_value(rows, 8) or 0
                if not product_qty:
                    message += f"\n\t - Không xác định số lượng tại hàng {rows + 1}."
                if product_code and product_qty:
                    value.append({
                        'product_code': product_code,
                        'product_qty': float(product_qty),
                        'po_name': po_name,
                        'description': description,
                        'so_number': so_number,
                        'order_date': order_date,
                        'delivery_date': delivery_date,
                        'price_unit': price_unit,
                        'discount': discount
                    })
            except Exception as e:
                message += f"\n\t - Lỗi định dạng mã sản phẩm/số lượng tại hàng {rows + 1}." \
                           f"\n\tChi tiết: {str(e)}"
            rows += 1
        if value and po_code:
            list_data.append([po_code, value])
        return list_data, message

    def _get_data_satra(self, txt_list):
        order_list = []
        for text in txt_list.split('|||||'):
            list_value = []
            location = text.split('Địa chỉ giao hàng: ')[1:][0].split('\nĐịa chỉ thanh toán:')[0].replace('\n', '')
            value = text.split('Ghi chú\n')[1:][0].split('\nTổng cộng')[0].split('\n')
            i = 1
            for item in value:
                item_vals = item.split(" ")
                x = 0
                i_in_vals = False
                for val in item_vals:
                    x += 1 if ',' in val else 0
                    if str(i) == val:
                        i_in_vals = True
                if x < 3:
                    i += 1 if i_in_vals else 0
                    continue
                while (True):
                    try:
                        index = item_vals.index(str(i))
                        code = item_vals[index + 1]
                        barcode = item_vals[index + 2]
                        quantity = 0
                        for val in item_vals:
                            if ',' in val:
                                try:
                                    quantity = float(val.replace(',', '.'))
                                except Exception:
                                    raise UserError("Hệ thống không đọc được số lượng trong file import")
                                break
                        if quantity > 0:
                            list_value.append((code, barcode, quantity))
                        i += 1
                    except Exception:
                        break
            order_list.append([location, list_value])
        return order_list

    def _get_data_brg(self, txt_list):
        if not txt_list:
            return '', [], " File là bản scan, Hệ thống không thể đọc!"
        value = []
        message = ''
        po_code = ''
        po_name = ''

        def isfloat(num):
            try:
                float(num)
                return True
            except ValueError:
                return False

        for text in txt_list.split('|||||'):
            # doc file bi loi font
            po_code = text.split('\nTæng gi¸ trÞ ®¬n hµng')[1:][0].split('|||')[0].replace('\n', '')
            po_name = text.split('|||')[1:][0].split('|||')[0].replace('\n', '')
            lines = text.split('®¬n nhËn\n')[1:][0].split('\nTæng gi¸ trÞ ®¬n hµng')[0].split('\n')
            for item in lines:
                vals = item.split(" ")
                item_vals = list(filter(lambda x: x.isdigit() or isfloat(x), vals))
                code = ''
                barcode = ''
                quantity = None
                if not item_vals:
                    continue
                if not item_vals[-3] and not item_vals[-2]:
                    message += f"\n\t - Thiếu mã sản phẩm/barcode tại STT {item_vals[0]}"
                else:
                    code = int(item_vals[-3]) or ''
                    barcode = int(item_vals[-2]) or ''
                if not item_vals[-1]:
                    message += f"\n\t - Thiếu số lượng tại STT {item_vals[0]}"
                else:
                    quantity = float(item_vals[-1])
                if code and barcode and quantity:
                    value.append((code, barcode, quantity, po_name))
        return po_code, value, message

    def _read_value_from_excel(self, attachment):
        return load_sheet(base64.b64decode(attachment.datas), attachment.name or '')

    def _read_value_from_import_SO_excel(self, attachment):
        return load_sheet(base64.b64decode(attachment), self.file_name or '')

    def download_import_SO_template(self):
        return {
            "type": "ir.actions.act_url",
            "url": '/sale_order_import/static/template/File_mau_import_excel.xlsx',
            "target": "_parent",
        }

    def download_import_SO_consignment_template(self):
        return {
            "type": "ir.actions.act_url",
            "url": '/sale_order_import/static/template/File_mau_import_consignment.xlsx',
            "target": "_parent",
        }

    def action_import_SO_excel(self):
        if not self.file_upload:
            raise UserError(_("File is not empty"))
        attachment = self.file_upload
        self.ensure_one()
        if self.state != 'draft':
            return
        if self.line_ids:
            self.line_ids.unlink()
        self.check_format_file_excel(self.file_name)
        lines = []
        messages = ''
        value_line, message_error = self._generate_order_from_attachment(attachment)
        if message_error:
            messages += message_error
        else:
            lines += value_line
        if messages:
            raise ValidationError(messages)
        self.line_ids = lines

    def check_format_file_excel(self, file_name):
        if not (file_name or '').lower().endswith(('.xls', '.xlsx')):
            self.file_upload = None
            self.file_name = None
            raise UserError(_("File must be 'xlsx' or 'xls'"))
        # TODO : đưa về 1 hàm khác vào các nút import
        if self.state != 'draft':
            raise UserError(_("State is not 'Draft'. Please Back to Draft for importing again"))
